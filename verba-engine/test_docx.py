import uuid
from docx import Document
from docx.shared import Inches
from docx_processor import DOCXProcessor

# 1. Create a mock DOCX
doc = Document()
doc.add_heading('Test Document', 0)
doc.add_paragraph('This is a normal paragraph.')

p = doc.add_paragraph('This is a bold paragraph. ')
p.runs[0].bold = True

table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'Header 1'
table.cell(0, 1).text = 'Header 2'
table.cell(1, 0).text = 'Data 1'
table.cell(1, 1).text = 'Data 2'

doc.save('test_fixture.docx')

# 2. Test the processor
with open('test_fixture.docx', 'rb') as f:
    docx_bytes = f.read()
    
# Mock credentials won't actually upload to supabase without correct env vars, 
# but it will test parsing logic
processor = DOCXProcessor(docx_bytes, user_id=str(uuid.uuid4()), document_id=str(uuid.uuid4()))
ast = processor.parse_to_json()

import json
print(json.dumps(ast, indent=2))
